"""Convert a skripsi chapter Markdown file into a formatted DOCX file."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

PROJECT_ROOT = Path(__file__).resolve().parents[1]

BOLD_ITALIC_PATTERN = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")


def add_runs_with_inline_formatting(paragraph, text: str) -> None:
    """Add text to a paragraph, honoring **bold** and *italic* markdown spans."""
    for part in BOLD_ITALIC_PATTERN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def set_base_style(document: Document) -> None:
    """Set default font to Times New Roman 12pt for the Normal style."""
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(12)


def convert(markdown_path: Path, docx_path: Path) -> None:
    """Convert one Markdown chapter file to DOCX."""
    document = Document()
    set_base_style(document)

    ordered_list_pattern = re.compile(r"^\d+\.\s+(.*)")
    lines = markdown_path.read_text(encoding="utf-8").splitlines()

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue

        if line.startswith("# "):
            heading = document.add_heading(level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs_with_inline_formatting(heading, line[2:])
            continue

        if line.startswith("## "):
            heading = document.add_heading(level=2)
            add_runs_with_inline_formatting(heading, line[3:])
            continue

        ordered_match = ordered_list_pattern.match(line)
        if ordered_match:
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs_with_inline_formatting(paragraph, ordered_match.group(1))
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Pt(36)
        add_runs_with_inline_formatting(paragraph, line)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(docx_path)


def parse_args() -> argparse.Namespace:
    """Parse CLI arguments."""
    parser = argparse.ArgumentParser(
        description="Convert a skripsi chapter Markdown file to DOCX."
    )
    parser.add_argument("--input", required=True, help="Path to the source .md file.")
    parser.add_argument(
        "--output",
        default=None,
        help="Path to the output .docx file (defaults to same name as input).",
    )
    return parser.parse_args()


def main() -> None:
    """CLI entrypoint."""
    args = parse_args()
    input_path = Path(args.input)
    output_path = Path(args.output) if args.output else input_path.with_suffix(".docx")
    convert(input_path, output_path)
    print(f"DOCX saved to: {output_path}")


if __name__ == "__main__":
    main()
