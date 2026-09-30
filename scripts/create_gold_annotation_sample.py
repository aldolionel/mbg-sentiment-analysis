"""Create a blind gold-annotation sample for human labeling.

Strata (from data/processed/mbg_labeled_sample_1000.csv):
- flagged: all rows flagged for review (adjudication_applied == True), taken in full.
- unflagged_negatif: all unflagged rows whose label is negatif, taken in full.
- unflagged_positif / unflagged_netral: random samples of the remaining unflagged rows.

The annotator file contains only an opaque item id and the text. The key file maps
item ids back to sample ids, strata, and the automatic labels, and must not be shown
to annotators. It is written to a *.private.csv path that is gitignored and can be
regenerated deterministically from the same seed and input.
"""

from __future__ import annotations

import argparse
import csv
import random
from pathlib import Path

LABELS = ["positif", "negatif", "netral"]


def norm(value: str) -> str:
    return value.strip().lower()


def build_sample(rows: list[dict], n_pos: int, n_neu: int, seed: int) -> list[dict]:
    rng = random.Random(seed)
    flagged = [r for r in rows if r["adjudication_applied"] == "True"]
    unflagged = [r for r in rows if r["adjudication_applied"] != "True"]
    by_label = {
        label: sorted(
            (r for r in unflagged if norm(r["label_after_adjudication"]) == label),
            key=lambda r: r["sample_id"],
        )
        for label in LABELS
    }
    strata = {
        "flagged": sorted(flagged, key=lambda r: r["sample_id"]),
        "unflagged_negatif": by_label["negatif"],
        "unflagged_positif": rng.sample(by_label["positif"], n_pos),
        "unflagged_netral": rng.sample(by_label["netral"], n_neu),
    }
    population = {
        "flagged": len(flagged),
        "unflagged_negatif": len(by_label["negatif"]),
        "unflagged_positif": len(by_label["positif"]),
        "unflagged_netral": len(by_label["netral"]),
    }
    selected = []
    for name, members in strata.items():
        for r in members:
            selected.append({**r, "stratum": name, "stratum_population": population[name]})
    rng.shuffle(selected)
    for i, r in enumerate(selected, start=1):
        r["item_id"] = f"G{i:03d}"
    return selected


def write_csv(path: Path, fieldnames: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path: Path, rows: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "annotasi"
    ws.append(["item_id", "clean_text", "annotator_label", "annotator_notes"])
    for r in rows:
        ws.append([r["item_id"], r["clean_text"], "", ""])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 90
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 40
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"positif,negatif,netral"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"C2:C{len(rows) + 1}")
    ws.freeze_panes = "A2"
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--input", default="data/processed/mbg_labeled_sample_1000.csv")
    parser.add_argument("--out-dir", default="data/processed/validation")
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--n-unflagged-positif", type=int, default=60)
    parser.add_argument("--n-unflagged-netral", type=int, default=60)
    args = parser.parse_args()

    with Path(args.input).open(newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    selected = build_sample(rows, args.n_unflagged_positif, args.n_unflagged_netral, args.seed)

    out = Path(args.out_dir)
    blind_fields = ["item_id", "clean_text", "annotator_label", "annotator_notes"]
    blind_rows = [{**r, "annotator_label": "", "annotator_notes": ""} for r in selected]
    write_csv(out / "gold_annotation_blind.csv", blind_fields, blind_rows)
    write_xlsx(out / "gold_annotation_blind.xlsx", selected)

    key_rows = [
        {
            "item_id": r["item_id"],
            "sample_id": r["sample_id"],
            "batch_id": r["batch_id"],
            "stratum": r["stratum"],
            "stratum_population": r["stratum_population"],
            "auto_label_initial": r["label_before_adjudication"],
            "auto_label_final": r["label_after_adjudication"],
        }
        for r in selected
    ]
    key_fields = list(key_rows[0].keys())
    write_csv(out / "gold_annotation_key.private.csv", key_fields, key_rows)

    counts: dict[str, int] = {}
    for r in selected:
        counts[r["stratum"]] = counts.get(r["stratum"], 0) + 1
    print(f"rows: {len(selected)}")
    for name, n in counts.items():
        print(f"  {name}: {n}")
    print(f"annotator files: {out / 'gold_annotation_blind.csv'}, {out / 'gold_annotation_blind.xlsx'}")
    print(f"key (do not share with annotators): {out / 'gold_annotation_key.private.csv'}")


if __name__ == "__main__":
    main()
