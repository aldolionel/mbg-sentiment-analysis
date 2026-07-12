"""Generate DOCX version of the revised methodology paper."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
MARKDOWN_PATH = PROJECT_ROOT / "paper" / "paper_mbg_sentiment_svm_revised_methodology.md"
TABLES_PATH = PROJECT_ROOT / "paper" / "paper_tables_revised_methodology.md"
DOCX_PATH = PROJECT_ROOT / "paper" / "paper_mbg_sentiment_svm_revised_methodology.docx"
REPORT_PATH = PROJECT_ROOT / "outputs" / "reports" / "22_revised_paper_report.md"

FIGURES = [
    PROJECT_ROOT / "outputs" / "figures" / "16_final_label_distribution.png",
    PROJECT_ROOT / "outputs" / "figures" / "16_model_comparison_macro_f1.png",
    PROJECT_ROOT / "outputs" / "figures" / "16_model_comparison_accuracy_vs_macro_f1.png",
    PROJECT_ROOT / "outputs" / "figures" / "16_linear_svm_improvement.png",
    PROJECT_ROOT / "outputs" / "figures" / "15_confusion_matrix_tuned_linearsvc.png",
    PROJECT_ROOT / "outputs" / "figures" / "16_negative_class_performance_tuned_svm.png",
]


def set_cell_shading(cell, fill: str) -> None:
    """Set cell background fill."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table) -> None:
    """Set table cell margins."""
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for side, value in [("top", "80"), ("bottom", "80"), ("start", "120"), ("end", "120")]:
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), value)
        element.set(qn("w:type"), "dxa")


def apply_styles(doc: Document) -> None:
    """Apply narrative proposal-style document styles using Times New Roman."""
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, RGBColor(46, 116, 181), 18, 10),
        ("Heading 2", 13, RGBColor(46, 116, 181), 12, 6),
        ("Heading 3", 12, RGBColor(31, 77, 120), 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_markdown_table(doc: Document, lines: list[str]) -> None:
    """Add a Markdown pipe table as a Word table."""
    rows = []
    for line in lines:
        stripped = line.strip()
        if not stripped.startswith("|") or set(stripped.replace("|", "").strip()) <= {"-", ":", " "}:
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        rows.append(cells)
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_cell_margins(table)
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = value
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9)
            if row_idx == 0:
                set_cell_shading(cell, "F4F6F9")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
    doc.add_paragraph()


def add_markdown_content(doc: Document, markdown: str, include_tables: bool = True) -> None:
    """Convert a practical subset of Markdown to DOCX content."""
    table_buffer: list[str] = []
    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        if line.startswith("|"):
            table_buffer.append(line)
            continue
        if table_buffer:
            if include_tables:
                add_markdown_table(doc, table_buffer)
            table_buffer = []
        if not line.strip():
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(line[2:].strip())
            run.font.name = "Times New Roman"
            run.font.size = Pt(16)
            run.bold = True
            run.font.color.rgb = RGBColor(11, 37, 69)
        elif line.startswith("## "):
            doc.add_paragraph(line[3:].strip(), style="Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:].strip(), style="Heading 2")
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.startswith("> "):
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(line[2:].strip())
            run.italic = True
        else:
            doc.add_paragraph(line)
    if table_buffer and include_tables:
        add_markdown_table(doc, table_buffer)


def add_figures(doc: Document) -> tuple[list[str], list[str]]:
    """Insert available figures and return inserted/missing lists."""
    inserted = []
    missing = []
    doc.add_page_break()
    doc.add_paragraph("Lampiran Gambar", style="Heading 1")
    for index, figure in enumerate(FIGURES, start=1):
        if not figure.exists():
            missing.append(str(figure))
            continue
        caption = doc.add_paragraph(f"Gambar {index}. {figure.name}")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in caption.runs:
            run.bold = True
        doc.add_picture(str(figure), width=Inches(5.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        inserted.append(str(figure))
    return inserted, missing


def create_docx() -> tuple[list[str], list[str]]:
    """Create revised paper DOCX."""
    doc = Document()
    apply_styles(doc)
    add_markdown_content(doc, MARKDOWN_PATH.read_text(encoding="utf-8"), include_tables=False)
    doc.add_page_break()
    doc.add_paragraph("Lampiran Tabel", style="Heading 1")
    add_markdown_content(doc, TABLES_PATH.read_text(encoding="utf-8"), include_tables=True)
    inserted, missing = add_figures(doc)
    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)
    return inserted, missing


def write_report(inserted: list[str], missing: list[str]) -> None:
    """Write revised paper generation report."""
    lines = [
        "# 22 Revised Paper Report",
        "",
        "## Files Created",
        f"- `paper/paper_mbg_sentiment_svm_revised_methodology.md`",
        f"- `paper/paper_tables_revised_methodology.md`",
        f"- `paper/paper_figures_revised_methodology.md`",
        f"- `paper/paper_mbg_sentiment_svm_revised_methodology.docx`",
        f"- `outputs/reports/22_revised_paper_report.md`",
        "",
        "## Major Changes from Previous Paper",
        "- Reframed the dataset as a secondary dataset associated with Sultoni et al. (2025).",
        "- Removed any claim that the student/researcher crawled the data.",
        "- Reframed contribution as methodology/evaluation rather than platform or SVM novelty.",
        "- Added cautious wording about CV evidence and test-set model selection.",
        "",
        "## Methodology Hardening Items Included",
        "- Dataset provenance subsection.",
        "- AI-assisted labeling and codebook limitation.",
        "- Class-weight ablation result.",
        "- CV model selection caution.",
        "- Near-duplicate train/test check result.",
        "- Negative-class qualitative error analysis.",
        "- Data and code availability section.",
        "",
        "## Remaining Limitations",
        "- Dataset license is not explicitly identified from available local documents.",
        "- Labels are not pure manual gold standard.",
        "- Cohen's Kappa has not been computed because two independent human annotators have not completed validation.",
        "- Results describe the sampled labeled dataset, not universal public opinion.",
        "",
        "## Figures Inserted",
    ]
    lines.extend([f"- `{path}`" for path in inserted] or ["- none"])
    lines.extend(["", "## Figures Missing"])
    lines.extend([f"- `{path}`" for path in missing] or ["- none"])
    lines.extend(
        [
            "",
            "## Readiness Status",
            "PASS: Revised paper Markdown, tables, figure guide, and DOCX were generated. DOCX still requires visual QA/render review before final submission.",
        ]
    )
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    """CLI entrypoint."""
    inserted, missing = create_docx()
    write_report(inserted, missing)
    print(f"Saved DOCX: {DOCX_PATH}")
    print(f"Saved report: {REPORT_PATH}")
    print(f"Figures inserted: {len(inserted)}")
    print(f"Figures missing: {len(missing)}")


if __name__ == "__main__":
    main()
