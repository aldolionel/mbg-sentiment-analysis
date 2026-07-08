"""Utilities for auditing raw sentiment-analysis datasets."""

from __future__ import annotations

import hashlib
import re
from collections import Counter
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xls"}
TEXT_COLUMN_KEYWORDS = {
    "komentar",
    "comment",
    "comments",
    "text",
    "tweet",
    "full_text",
    "content",
    "caption",
    "description",
    "clean_text",
}
LABEL_COLUMN_KEYWORDS = {
    "label",
    "sentiment",
    "sentimen",
    "polarity",
    "polaritas",
    "kelas",
    "category",
    "kategori",
}
NON_TEXT_METADATA_KEYWORDS = {
    "created",
    "date",
    "time",
    "id",
    "url",
    "link",
    "screen_name",
    "username",
    "user",
    "name",
    "location",
    "lang",
}
TOKEN_PATTERN = re.compile(r"[a-zA-ZÀ-ÿ]+")


def list_raw_data_files(raw_dir: str | Path = "data/raw") -> list[Path]:
    """List supported tabular raw data files.

    Args:
        raw_dir: Directory containing raw datasets.

    Returns:
        Sorted paths for CSV and Excel files.
    """
    raw_path = Path(raw_dir)
    if not raw_path.exists():
        return []

    return sorted(
        path
        for path in raw_path.iterdir()
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )


def excel_sheet_names(path: str | Path) -> list[str]:
    """Return sheet names for an Excel file.

    Args:
        path: Path to an Excel file.

    Returns:
        List of sheet names.

    Raises:
        FileNotFoundError: If the file does not exist.
        ValueError: If the file is not an Excel file.
    """
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")
    if file_path.suffix.lower() not in {".xlsx", ".xls"}:
        raise ValueError(f"Not an Excel file: {file_path}")

    return list(pd.ExcelFile(file_path).sheet_names)


def load_table(
    path: str | Path,
    sheet_name: str | int | None = None,
) -> pd.DataFrame:
    """Load a CSV or Excel table into a DataFrame.

    Args:
        path: Path to a CSV or Excel file.
        sheet_name: Excel sheet name or index. If None, the first sheet is used.

    Returns:
        Loaded DataFrame.

    Raises:
        FileNotFoundError: If the file does not exist.
        ValueError: If the file extension is unsupported.
    """
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        return pd.read_csv(file_path)
    if suffix == ".xlsx":
        selected_sheet = 0 if sheet_name is None else sheet_name
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            if isinstance(selected_sheet, int):
                worksheet = workbook.worksheets[selected_sheet]
            else:
                worksheet = workbook[str(selected_sheet)]

            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                return pd.DataFrame()

            columns = [
                str(column) if column is not None else f"unnamed_{index}"
                for index, column in enumerate(header)
            ]
            return pd.DataFrame(rows, columns=columns)
        finally:
            workbook.close()

    if suffix == ".xls":
        selected_sheet = 0 if sheet_name is None else sheet_name
        return pd.read_excel(file_path, sheet_name=selected_sheet)

    raise ValueError(
        f"Unsupported file type '{suffix}'. "
        f"Supported types: {sorted(SUPPORTED_EXTENSIONS)}"
    )


def profile_dataframe(df: pd.DataFrame) -> dict[str, Any]:
    """Build a concise profile for a DataFrame.

    Args:
        df: DataFrame to profile.

    Returns:
        Dictionary containing shape, missingness, duplicates, memory, and sample.
    """
    missing_values = df.isna().sum()
    missing_rate = df.isna().mean()
    duplicate_rows = int(df.duplicated().sum())
    n_rows = int(len(df))

    return {
        "n_rows": n_rows,
        "n_columns": int(df.shape[1]),
        "column_names": list(df.columns),
        "dtypes": {col: str(dtype) for col, dtype in df.dtypes.items()},
        "missing_values": {
            col: int(value) for col, value in missing_values.items()
        },
        "missing_rate": {
            col: float(value) for col, value in missing_rate.items()
        },
        "duplicate_rows": duplicate_rows,
        "duplicate_rate": (
            float(duplicate_rows / n_rows) if n_rows > 0 else 0.0
        ),
        "memory_mb": float(
            df.memory_usage(deep=True).sum() / (1024**2)
        ),
        "sample_rows": df.head(5).where(pd.notna(df.head(5)), None).to_dict(
            orient="records"
        ),
    }


def infer_text_columns(df: pd.DataFrame) -> list[str]:
    """Infer likely text/comment columns.

    Args:
        df: DataFrame to inspect.

    Returns:
        Ordered list of likely text columns.
    """
    keyword_matches: list[str] = []
    length_matches: list[str] = []
    for col in df.columns:
        normalized = str(col).strip().lower()
        if normalized in TEXT_COLUMN_KEYWORDS:
            keyword_matches.append(str(col))
            continue

        if any(keyword in normalized for keyword in TEXT_COLUMN_KEYWORDS):
            keyword_matches.append(str(col))
            continue

        if any(keyword in normalized for keyword in NON_TEXT_METADATA_KEYWORDS):
            continue

        if pd.api.types.is_object_dtype(df[col]) or pd.api.types.is_string_dtype(
            df[col]
        ):
            text_values = df[col].dropna().astype(str).str.strip()
            if text_values.empty:
                continue
            avg_length = text_values.str.len().mean()
            unique_ratio = text_values.nunique() / len(text_values)
            if avg_length >= 25 and unique_ratio >= 0.2:
                length_matches.append(str(col))

    return list(dict.fromkeys(keyword_matches + length_matches))


def infer_label_columns(df: pd.DataFrame) -> list[str]:
    """Infer likely sentiment label columns.

    Args:
        df: DataFrame to inspect.

    Returns:
        Ordered list of likely label columns.
    """
    inferred: list[str] = []
    for col in df.columns:
        normalized = str(col).strip().lower()
        if normalized in LABEL_COLUMN_KEYWORDS:
            inferred.append(str(col))
            continue

        if any(keyword in normalized for keyword in LABEL_COLUMN_KEYWORDS):
            inferred.append(str(col))

    return list(dict.fromkeys(inferred))


def _basic_tokens(text_values: pd.Series) -> list[str]:
    """Tokenize text with a simple alphabetic regex."""
    tokens: list[str] = []
    for value in text_values.dropna().astype(str):
        tokens.extend(TOKEN_PATTERN.findall(value.lower()))
    return tokens


def text_column_stats(df: pd.DataFrame, text_col: str) -> dict[str, Any]:
    """Compute basic quality and length statistics for a text column.

    Args:
        df: DataFrame containing the text column.
        text_col: Text column name.

    Returns:
        Dictionary with text completeness, duplication, length, and terms.

    Raises:
        KeyError: If the text column is not present.
    """
    if text_col not in df.columns:
        raise KeyError(f"Text column not found: {text_col}")

    series = df[text_col]
    non_null = series.dropna().astype(str)
    stripped = non_null.str.strip()
    non_empty = stripped[stripped != ""]
    lengths = non_empty.str.len()
    word_counts = non_empty.apply(
        lambda value: len(TOKEN_PATTERN.findall(value.lower()))
    )
    tokens = _basic_tokens(non_empty)
    top_terms = Counter(tokens).most_common(20)

    return {
        "non_null_count": int(series.notna().sum()),
        "empty_string_count": int((stripped == "").sum()),
        "duplicate_text_count": int(non_empty.duplicated().sum()),
        "avg_length_chars": (
            float(lengths.mean()) if not lengths.empty else 0.0
        ),
        "median_length_chars": (
            float(lengths.median()) if not lengths.empty else 0.0
        ),
        "min_length_chars": int(lengths.min()) if not lengths.empty else 0,
        "max_length_chars": int(lengths.max()) if not lengths.empty else 0,
        "avg_word_count": (
            float(word_counts.mean()) if not word_counts.empty else 0.0
        ),
        "median_word_count": (
            float(word_counts.median()) if not word_counts.empty else 0.0
        ),
        "top_20_terms_basic": [
            {"term": term, "count": int(count)}
            for term, count in top_terms
        ],
    }


def label_distribution(df: pd.DataFrame, label_col: str) -> pd.DataFrame:
    """Calculate label counts and percentages.

    Args:
        df: DataFrame containing labels.
        label_col: Label column name.

    Returns:
        DataFrame with count and percentage per label.

    Raises:
        KeyError: If the label column is not present.
    """
    if label_col not in df.columns:
        raise KeyError(f"Label column not found: {label_col}")

    counts = df[label_col].fillna("<missing>").astype(str).value_counts()
    percentages = counts / counts.sum() * 100 if counts.sum() else counts
    return pd.DataFrame(
        {
            "label": counts.index,
            "count": counts.values.astype(int),
            "percentage": percentages.values,
        }
    )


def file_sha256(path: str | Path, first_n_chars: int = 12) -> str:
    """Compute a SHA-256 file fingerprint.

    Args:
        path: File path.
        first_n_chars: Number of leading hash characters to return.

    Returns:
        SHA-256 hash prefix.

    Raises:
        FileNotFoundError: If the file does not exist.
    """
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    sha256 = hashlib.sha256()
    with file_path.open("rb") as file_obj:
        for chunk in iter(lambda: file_obj.read(1024 * 1024), b""):
            sha256.update(chunk)
    return sha256.hexdigest()[:first_n_chars]


def audit_raw_datasets(raw_dir: str | Path = "data/raw") -> dict[str, Any]:
    """Audit all supported raw datasets in a directory.

    Excel files are audited sheet by sheet. CSV files are represented with a
    single pseudo-sheet named ``<csv>``.

    Args:
        raw_dir: Directory containing raw datasets.

    Returns:
        Nested audit dictionary by file and sheet.
    """
    files = list_raw_data_files(raw_dir)
    audit: dict[str, Any] = {
        "raw_dir": str(Path(raw_dir)),
        "files_found": [str(path) for path in files],
        "datasets": [],
    }

    for file_path in files:
        file_entry: dict[str, Any] = {
            "path": str(file_path),
            "name": file_path.name,
            "extension": file_path.suffix.lower(),
            "sha256": file_sha256(file_path),
            "sheets": {},
        }

        if file_path.suffix.lower() in {".xlsx", ".xls"}:
            sheet_names = excel_sheet_names(file_path)
        else:
            sheet_names = ["<csv>"]

        for sheet in sheet_names:
            try:
                df = load_table(
                    file_path,
                    sheet_name=None if sheet == "<csv>" else sheet,
                )
                file_entry["sheets"][sheet] = {
                    "profile": profile_dataframe(df),
                    "inferred_text_columns": infer_text_columns(df),
                    "inferred_label_columns": infer_label_columns(df),
                    "text_stats": {
                        text_col: text_column_stats(df, text_col)
                        for text_col in infer_text_columns(df)
                    },
                    "label_distributions": {
                        label_col: label_distribution(df, label_col).to_dict(
                            orient="records"
                        )
                        for label_col in infer_label_columns(df)
                    },
                }
            except Exception as exc:  # pragma: no cover - defensive report path
                file_entry["sheets"][sheet] = {"error": str(exc)}

        audit["datasets"].append(file_entry)

    return audit
